import io
import os
import re
import openpyxl
from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, jsonify, g
from sqlalchemy import or_, and_, func
from app.extensions import db
from app.models import Urun, StokHareketi
from app.auth import login_required

stok_bp = Blueprint("stok", __name__)


def _temiz_float(val, default=0.0):
    if val is None:
        return default
    s = str(val).strip().replace("TL", "").replace("$", "").replace(" ", "")
    if not s:
        return default
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return default


def akilli_urun_sorgusu_uygula(sorgu, q):
    """
    Kullanıcı 'tiandy TC32GN', 'tc34xn', 'TC-C32GN', 'platoon 1866' veya 'DS-FG4' yazdığında
    kelime kelime ve tire/boşluk farklarını (TC-C32GN <-> TC32GN) tolere ederek filtreler.
    """
    if not q or not q.strip():
        return sorgu

    kelimeler = [k.strip() for k in q.strip().split() if k.strip()]
    norm_kod = func.replace(func.replace(func.lower(func.coalesce(Urun.kod, "")), "-", ""), " ", "")
    norm_ad = func.replace(func.lower(func.coalesce(Urun.ad, "")), "-", "")

    for kelime in kelimeler:
        k_lower = kelime.lower()
        k_temiz = k_lower.replace("-", "").replace(" ", "")
        arama = f"%{kelime}%"
        arama_temiz = f"%{k_temiz}%"

        kosullar = [
            Urun.barkod.ilike(arama),
            Urun.kod.ilike(arama),
            Urun.ad.ilike(arama),
            Urun.marka.ilike(arama),
            Urun.grup.ilike(arama),
            norm_kod.like(arama_temiz),
            norm_ad.like(arama_temiz),
        ]

        # Örn: Kullanıcı 'TC32GN' veya 'tc34xn' yazdıysa ama kod 'TC-C32GN' / 'TC-C34XN' / 'TC-H333N' ise:
        m = re.match(r"^tc([0-9a-z]+)$", k_temiz)
        if m:
            sonek = m.group(1)
            kosullar.append(norm_kod.like(f"%tc%{sonek}%"))
            kosullar.append(Urun.kod.ilike(f"%{sonek}%"))

        sorgu = sorgu.filter(or_(*kosullar))

    return sorgu


def _excel_calisma_sayfasini_aktar(ws, kullanici_id=None, guncelle=True):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0, 0

    baslik_satiri = [str(h or "").strip().lower() for h in rows[0]]

    def sutun_bul(adaylar, varsayilan=None, haric_kelimeler=None):
        haric_kelimeler = haric_kelimeler or []
        # Önce tam eşleşme ara
        for a in adaylar:
            for idx, h in enumerate(baslik_satiri):
                if any(hk in h for hk in haric_kelimeler):
                    continue
                if h == a:
                    return idx
        # Sonra alt-metin eşleşmesi ara
        for a in adaylar:
            for idx, h in enumerate(baslik_satiri):
                if any(hk in h for hk in haric_kelimeler):
                    continue
                if a in h:
                    return idx
        return varsayilan

    idx_barkod = sutun_bul(["barkod", "barcode"], 0)
    idx_grup = sutun_bul(["grup", "group"], 1)
    idx_ad = sutun_bul(["ürün adı", "urun adi", "ürün", "product", "katagori", "kategori", "ad"], 2)
    idx_marka = sutun_bul(["marka", "brand"], 3)
    # 'barkod' sütununun 'kod' kelimesiyle eşleşmesini engellemek için haric_kelimeler=['barkod']
    idx_kod = sutun_bul(["model", "stok kodu", "ürün kodu", "kod", "code"], 4, haric_kelimeler=["barkod"])
    idx_adet = sutun_bul(["adet", "miktar", "stok", "qty"], 5)
    idx_alis_tl = sutun_bul(["alış fiyatı tl", "alis fiyati tl", "tl", "maliyet"], 6)
    idx_alis_usd = sutun_bul(["dolar", "usd", "alış fiyatı dolar"], 7)
    idx_satis = sutun_bul(["satış", "satis", "fiyat"], 8, haric_kelimeler=["alış", "alis"])

    eklenen = 0
    guncellenen = 0

    for satir in rows[1:]:
        if not satir or all(c is None or str(c).strip() == "" for c in satir):
            continue

        def hucre(i):
            if i is not None and i < len(satir) and satir[i] is not None:
                return str(satir[i]).strip()
            return ""

        barkod = hucre(idx_barkod)
        if barkod.endswith(".0"):
            barkod = barkod[:-2]
        grup = hucre(idx_grup) or "Genel"
        ad = hucre(idx_ad)
        marka = hucre(idx_marka)
        kod = hucre(idx_kod)
        if kod.endswith(".0"):
            kod = kod[:-2]

        # Ara başlık satırlarını atla (Ör: Row 282 'KAMERALAR' / '2026-07-07')
        if ad.startswith("2026-") or (barkod.upper() == "KAMERALAR" and not kod):
            continue

        if not ad:
            ad = f"{marka} {kod}".strip()
        if not ad:
            continue

        stok_miktari = _temiz_float(hucre(idx_adet), 0.0)
        alis_tl = _temiz_float(hucre(idx_alis_tl), 0.0)
        alis_usd = _temiz_float(hucre(idx_alis_usd), 0.0)
        satis = _temiz_float(hucre(idx_satis), 0.0)

        mevcut = None
        if barkod and kod:
            mevcut = Urun.query.filter_by(barkod=barkod, kod=kod).first()
            if not mevcut:
                # Önceki hatalı aktarımda kod sütununa barkod yazılmış olan kaydı yakala ve düzelt
                mevcut = Urun.query.filter_by(barkod=barkod, kod=barkod, ad=ad).first()
        elif kod and marka:
            mevcut = Urun.query.filter_by(kod=kod, marka=marka).first()
        elif barkod:
            mevcut = Urun.query.filter_by(barkod=barkod, ad=ad).first()

        if mevcut and guncelle:
            onceki = mevcut.stok_miktari
            mevcut.ad = ad or mevcut.ad
            mevcut.grup = grup or mevcut.grup
            mevcut.marka = marka or mevcut.marka
            mevcut.kod = kod or None
            mevcut.stok_miktari = stok_miktari
            if alis_tl > 0:
                mevcut.alis_fiyati_tl = alis_tl
            if alis_usd > 0:
                mevcut.alis_fiyati_usd = alis_usd
            if satis > 0:
                mevcut.satis_fiyati = satis
            mevcut.aktif = True
            if onceki != stok_miktari:
                db.session.add(
                    StokHareketi(
                        urun_id=mevcut.id,
                        kullanici_id=kullanici_id,
                        hareket_tipi="sayim",
                        miktar_degisimi=round(stok_miktari - onceki, 3),
                        onceki_stok=onceki,
                        sonraki_stok=stok_miktari,
                        aciklama="Excel toplu güncelleme",
                    )
                )
            guncellenen += 1
        else:
            yeni = Urun(
                barkod=barkod or None,
                kod=kod or None,
                ad=ad,
                marka=marka,
                grup=grup,
                birim="Adet",
                alis_fiyati_tl=alis_tl,
                alis_fiyati_usd=alis_usd,
                satis_fiyati=satis,
                stok_miktari=stok_miktari,
                kritik_stok=2.0,
            )
            db.session.add(yeni)
            db.session.flush()
            if stok_miktari != 0:
                db.session.add(
                    StokHareketi(
                        urun_id=yeni.id,
                        kullanici_id=kullanici_id,
                        hareket_tipi="manuel_giris",
                        miktar_degisimi=stok_miktari,
                        onceki_stok=0.0,
                        sonraki_stok=stok_miktari,
                        aciklama="Excel ilk stok aktarımı",
                    )
                )
            eklenen += 1

    db.session.commit()
    return eklenen, guncellenen


@stok_bp.route("/urunler", methods=["GET", "POST"])
@login_required
def urunler_listesi():
    if request.method == "POST":
        if not (g.user.is_patron or g.user.yetki_stok_duzenle):
            flash("Stok/ürün düzenleme yetkiniz yok.", "danger")
            return redirect(url_for("stok.urunler_listesi"))

        form_islem = request.form.get("form_islem", "ekle")
        if form_islem == "ekle":
            barkod = request.form.get("barkod", "").strip()
            kod = request.form.get("kod", "").strip()
            ad = request.form.get("ad", "").strip()
            marka = request.form.get("marka", "").strip()
            grup = request.form.get("grup", "").strip() or "Genel"
            birim = request.form.get("birim", "Adet").strip()
            stok_miktari = _temiz_float(request.form.get("stok_miktari"), 0.0)
            kritik_stok = _temiz_float(request.form.get("kritik_stok"), 2.0)
            alis_tl = _temiz_float(request.form.get("alis_fiyati_tl"), 0.0)
            alis_usd = _temiz_float(request.form.get("alis_fiyati_usd"), 0.0)
            satis = _temiz_float(request.form.get("satis_fiyati"), 0.0)
            satis_2 = _temiz_float(request.form.get("satis_fiyati_2"), 0.0)

            if not ad:
                flash("Ürün adı boş bırakılamaz!", "danger")
            else:
                yeni = Urun(
                    barkod=barkod or None,
                    kod=kod or None,
                    ad=ad,
                    marka=marka,
                    grup=grup,
                    birim=birim,
                    stok_miktari=stok_miktari,
                    kritik_stok=kritik_stok,
                    alis_fiyati_tl=alis_tl,
                    alis_fiyati_usd=alis_usd,
                    satis_fiyati=satis,
                    satis_fiyati_2=satis_2,
                )
                db.session.add(yeni)
                db.session.flush()
                if stok_miktari != 0:
                    db.session.add(
                        StokHareketi(
                            urun_id=yeni.id,
                            kullanici_id=g.user.id,
                            hareket_tipi="manuel_giris",
                            miktar_degisimi=stok_miktari,
                            onceki_stok=0.0,
                            sonraki_stok=stok_miktari,
                            aciklama="Yeni ürün kartı açılışı",
                        )
                    )
                db.session.commit()
                flash(f"'{ad}' ({kod or 'Kodsuz'}) stoğa eklendi.", "success")

        elif form_islem == "duzenle":
            urun_id = int(request.form.get("urun_id", 0))
            u = db.session.get(Urun, urun_id)
            if u:
                onceki_stok = u.stok_miktari
                u.barkod = request.form.get("barkod", "").strip() or None
                u.kod = request.form.get("kod", "").strip() or None
                u.ad = request.form.get("ad", u.ad).strip()
                u.marka = request.form.get("marka", "").strip()
                u.grup = request.form.get("grup", "Genel").strip() or "Genel"
                u.birim = request.form.get("birim", "Adet").strip()
                yeni_stok = _temiz_float(request.form.get("stok_miktari"), u.stok_miktari)
                u.stok_miktari = yeni_stok
                u.kritik_stok = _temiz_float(request.form.get("kritik_stok"), u.kritik_stok)
                u.satis_fiyati = _temiz_float(request.form.get("satis_fiyati"), u.satis_fiyati)
                u.satis_fiyati_2 = _temiz_float(request.form.get("satis_fiyati_2"), u.satis_fiyati_2)

                if g.user.is_patron or g.user.yetki_maliyet_gor:
                    u.alis_fiyati_tl = _temiz_float(request.form.get("alis_fiyati_tl"), u.alis_fiyati_tl)
                    u.alis_fiyati_usd = _temiz_float(request.form.get("alis_fiyati_usd"), u.alis_fiyati_usd)

                if onceki_stok != yeni_stok:
                    db.session.add(
                        StokHareketi(
                            urun_id=u.id,
                            kullanici_id=g.user.id,
                            hareket_tipi="sayim",
                            miktar_degisimi=round(yeni_stok - onceki_stok, 3),
                            onceki_stok=onceki_stok,
                            sonraki_stok=yeni_stok,
                            aciklama="Ürün kartından stok düzenleme",
                        )
                    )
                db.session.commit()
                flash(f"'{u.ad}' ({u.kod or ''}) güncellendi.", "success")

        return redirect(url_for("stok.urunler_listesi", q=request.args.get("q", ""), grup=request.args.get("grup", "")))

    q = request.args.get("q", "").strip()
    secili_grup = request.args.get("grup", "").strip()
    durum = request.args.get("durum", "").strip()

    sorgu = Urun.query.filter_by(aktif=True)
    if secili_grup:
        sorgu = sorgu.filter(Urun.grup == secili_grup)
    if durum == "kritik":
        sorgu = sorgu.filter(Urun.stok_miktari <= Urun.kritik_stok)
    elif durum == "tukenen":
        sorgu = sorgu.filter(Urun.stok_miktari <= 0)

    if q:
        sorgu = akilli_urun_sorgusu_uygula(sorgu, q)

    urunler = sorgu.order_by(Urun.grup.asc(), Urun.ad.asc(), Urun.kod.asc()).all()
    gruplar = [
        r[0]
        for r in db.session.query(Urun.grup)
        .filter(Urun.aktif == True, Urun.grup != None, Urun.grup != "")
        .distinct()
        .order_by(Urun.grup.asc())
        .all()
    ]

    return render_template(
        "urunler.html",
        urunler=urunler,
        gruplar=gruplar,
        q=q,
        secili_grup=secili_grup,
        durum=durum,
    )


@stok_bp.route("/urunler/<int:id>/sil", methods=["POST"])
@login_required
def urun_sil(id):
    if not (g.user.is_patron or g.user.yetki_urun_sil):
        flash("Ürün silme yetkiniz bulunmuyor.", "danger")
        return redirect(url_for("stok.urunler_listesi"))

    u = db.session.get(Urun, id)
    if u:
        u.aktif = False
        db.session.commit()
        flash(f"'{u.ad}' silindi.", "warning")
    return redirect(url_for("stok.urunler_listesi"))


@stok_bp.route("/urunler/toplu-fiyatlandir", methods=["POST"])
@login_required
def toplu_fiyatlandir():
    if not g.user.is_patron:
        flash("Toplu fiyatlandırma yetkisi sadece Patron'a aittir.", "danger")
        return redirect(url_for("stok.urunler_listesi"))

    kar_yuzdesi = _temiz_float(request.form.get("kar_yuzdesi"), 35.0)
    kapsam = request.form.get("kapsam", "sadece_sifir")  # 'sadece_sifir' veya 'tumu'
    secili_grup = request.form.get("grup", "").strip()
    dolar_kuru = g.ayarlar.dolar_kuru if g.ayarlar else 35.0

    sorgu = Urun.query.filter_by(aktif=True)
    if secili_grup:
        sorgu = sorgu.filter(Urun.grup == secili_grup)
    if kapsam == "sadece_sifir":
        sorgu = sorgu.filter(or_(Urun.satis_fiyati == None, Urun.satis_fiyati <= 0))

    guncellenen = 0
    for u in sorgu.all():
        maliyet_tl = u.aktif_maliyet_tl(dolar_kuru)
        if maliyet_tl > 0:
            yeni_satis = round(maliyet_tl * (1.0 + (kar_yuzdesi / 100.0)), 2)
            u.satis_fiyati = yeni_satis
            guncellenen += 1

    db.session.commit()
    flash(
        f"Dolar kuru ({dolar_kuru} ₺) ve %{kar_yuzdesi:g} kâr marjı ile {guncellenen} ürünün satış fiyatı otomatik hesaplandı!",
        "success",
    )
    return redirect(url_for("stok.urunler_listesi", grup=secili_grup))


@stok_bp.route("/api/urun-hareketleri/<int:id>")
@login_required
def api_urun_hareketleri(id):
    u = db.session.get(Urun, id)
    if not u:
        return jsonify({"basarili": False}), 404

    from app.models import Islem
    hareketler = (
        StokHareketi.query.filter_by(urun_id=u.id)
        .order_by(StokHareketi.tarih.desc())
        .limit(40)
        .all()
    )
    liste = []
    for h in hareketler:
        islem = db.session.get(Islem, h.islem_id) if h.islem_id else None
        musteri_ad = ""
        montaj_yeri = ""
        fis_no = ""
        if islem:
            fis_no = islem.fis_no
            if islem.musteri:
                musteri_ad = islem.musteri.ad_soyad
            elif islem.tedarikci:
                musteri_ad = islem.tedarikci.firma_adi
            montaj_yeri = islem.montaj_adresi or islem.montaj_basligi or ""

        liste.append({
            "tarih": h.tarih.strftime("%d.%m.%Y %H:%M"),
            "tip": h.hareket_tipi.upper(),
            "degisim": h.miktar_degisimi,
            "sonraki_stok": h.sonraki_stok,
            "personel": h.kullanici.ad_soyad if h.kullanici else "-",
            "aciklama": h.aciklama or "",
            "musteri": musteri_ad,
            "montaj_yeri": montaj_yeri,
            "fis_no": fis_no,
            "islem_id": h.islem_id,
        })

    return jsonify({
        "basarili": True,
        "urun_ad": f"{u.ad} [{u.kod or ''}]",
        "hareketler": liste,
    })


@stok_bp.route("/urunler/excel", methods=["GET", "POST"])
@login_required
def excel_islemleri():
    if not (g.user.is_patron or g.user.yetki_stok_duzenle):
        flash("Excel aktarım yetkiniz bulunmuyor.", "danger")
        return redirect(url_for("stok.urunler_listesi"))

    varsayilan_excel_yolu = r"C:\Users\senel\Desktop\dükkan\excel stok\stok2 .xlsx"
    hazir_dosya_var = os.path.exists(varsayilan_excel_yolu)

    if request.method == "POST":
        kaynak = request.form.get("kaynak", "dosya")
        guncelle = "uzerine_yaz" in request.form

        try:
            if kaynak == "hazir_stok2" and hazir_dosya_var:
                wb = openpyxl.load_workbook(varsayilan_excel_yolu, data_only=True)
            else:
                dosya = request.files.get("excel_dosyasi")
                if not dosya or not dosya.filename:
                    flash("Lütfen bir Excel (.xlsx) dosyası seçin.", "danger")
                    return redirect(url_for("stok.excel_islemleri"))
                wb = openpyxl.load_workbook(dosya, data_only=True)

            ws = wb.active
            eklenen, guncellenen = _excel_calisma_sayfasini_aktar(
                ws, kullanici_id=g.user.id, guncelle=guncelle
            )
            flash(
                f"Excel aktarımı tamamlandı! {eklenen} yeni ürün eklendi, {guncellenen} mevcut ürün güncellendi.",
                "success",
            )
            return redirect(url_for("stok.urunler_listesi"))
        except Exception as e:
            flash(f"Excel okuma hatası: {e}", "danger")

    return render_template("excel.html", hazir_dosya_var=hazir_dosya_var, varsayilan_excel_yolu=varsayilan_excel_yolu)


@stok_bp.route("/urunler/excel-indir")
@login_required
def excel_indir():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Şeneller Stok"

    maliyet_izni = g.user.is_patron or g.user.yetki_maliyet_gor
    if maliyet_izni:
        ws.append([
            "Barkod",
            "grup",
            "Katagori",
            "Marka",
            "Model",
            "Adet",
            "Alış Fiyatı TL",
            "Alış Fiyatı DOLAR",
            "Satış Fiyatı",
        ])
    else:
        ws.append(["Barkod", "grup", "Katagori", "Marka", "Model", "Adet", "Satış Fiyatı"])

    for u in Urun.query.filter_by(aktif=True).order_by(Urun.grup.asc(), Urun.ad.asc()).all():
        if maliyet_izni:
            ws.append([
                u.barkod or "",
                u.grup or "",
                u.ad or "",
                u.marka or "",
                u.kod or "",
                u.stok_miktari,
                u.alis_fiyati_tl,
                u.alis_fiyati_usd,
                u.satis_fiyati,
            ])
        else:
            ws.append([
                u.barkod or "",
                u.grup or "",
                u.ad or "",
                u.marka or "",
                u.kod or "",
                u.stok_miktari,
                u.satis_fiyati,
            ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="Seneller_POS_Stok_Listesi.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@stok_bp.route("/sayim", methods=["GET", "POST"])
@login_required
def hizli_sayim():
    if request.method == "POST":
        veri = request.get_json(silent=True) or {}
        urun_id = veri.get("urun_id")
        u = db.session.get(Urun, int(urun_id)) if urun_id else None
        if not u:
            return jsonify({"basarili": False, "mesaj": "Ürün bulunamadı."}), 404

        onceki = u.stok_miktari
        yeni_stok = _temiz_float(veri.get("stok_miktari"), onceki)
        u.stok_miktari = yeni_stok
        if veri.get("satis_fiyati") is not None and (g.user.is_patron or g.user.yetki_fiyat_degistir):
            u.satis_fiyati = _temiz_float(veri.get("satis_fiyati"), u.satis_fiyati)
        if (g.user.is_patron or g.user.yetki_maliyet_gor) and veri.get("alis_fiyati_tl") is not None:
            u.alis_fiyati_tl = _temiz_float(veri.get("alis_fiyati_tl"), u.alis_fiyati_tl)

        if onceki != yeni_stok:
            db.session.add(
                StokHareketi(
                    urun_id=u.id,
                    kullanici_id=g.user.id,
                    hareket_tipi="sayim",
                    miktar_degisimi=round(yeni_stok - onceki, 3),
                    onceki_stok=onceki,
                    sonraki_stok=yeni_stok,
                    aciklama="Mobil/Barkodlu Hızlı Stok Sayımı",
                )
            )
        db.session.commit()
        return jsonify({
            "basarili": True,
            "mesaj": f"{u.ad} ({u.kod or ''}) güncellendi (Stok: {u.stok_miktari})",
            "stok_miktari": u.stok_miktari,
            "satis_fiyati": u.satis_fiyati,
        })

    son_hareketler = (
        StokHareketi.query.order_by(StokHareketi.tarih.desc()).limit(25).all()
    )
    return render_template("sayim.html", son_hareketler=son_hareketler)
